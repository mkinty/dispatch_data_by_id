from openpyxl import load_workbook

from dispatch_sna.services.dispatcher import ranger_commune, ranger_ids
from tests.fichiers import arbo, excel, log, textes


def test_ranger_ids_modifier(tmp_path):
    root, src_dir, dst_dir = arbo(tmp_path)
    st = ranger_ids(str(root), "74143", str(src_dir), ["74143_1"], "modifier", True, log)
    assert st == {"lignes": 1, "excel": 1, "ppt": 1, "err": 0}
    t = textes(dst_dir / "analyse" / "cas_analyse_74143_1.pptx")
    assert "Rem 1" in t and "Lien Street View : https://sv/1" in t
    assert "SNA nombre de logement: 2" in t        # c'est bien l'ancien PPT modifié
    assert not src_dir.exists()                     # tout traité → dossier supprimé


def test_ranger_ids_remplacer_et_conserve_dossier_si_autres_ids(tmp_path):
    root, src_dir, dst_dir = arbo(tmp_path)
    excel(src_dir / "audit_74143.xlsx", [["74143_1", "A1", "Rem 1", "u", None],
                                         ["74143_2", "A2", "Rem 2", "u", None]])
    st = ranger_ids(str(root), "74143", str(src_dir), ["74143_1"], "remplacer", True, log)
    assert st["err"] == 0 and st["ppt"] == 1
    assert "Justification IA" in textes(dst_dir / "analyse" / "cas_analyse_74143_1.pptx")
    assert src_dir.exists()                         # 74143_2 non traité → gardé
    ws = load_workbook(dst_dir / "audit_74143.xlsx").active
    assert ws["C3"].value == "vierge"


def test_ranger_ids_sans_option_ppt(tmp_path):
    root, src_dir, dst_dir = arbo(tmp_path)
    avant = open(dst_dir / "analyse" / "cas_analyse_74143_1.pptx", "rb").read()
    st = ranger_ids(str(root), "74143", str(src_dir), ["74143_1"], None, False, log)
    assert st == {"lignes": 1, "excel": 1, "ppt": 0, "err": 0}
    assert open(dst_dir / "analyse" / "cas_analyse_74143_1.pptx", "rb").read() == avant


def test_ranger_ids_aucune_colonne_mais_ppt_modifie(tmp_path):
    root, src_dir, dst_dir = arbo(tmp_path)
    st = ranger_ids(str(root), "74143", str(src_dir), ["74143_1"], "modifier", True, log, colonnes=[])
    assert st == {"lignes": 0, "excel": 0, "ppt": 1, "err": 0}
    assert load_workbook(dst_dir / "audit_74143.xlsx").active["C2"].value == "vierge"
    assert "Rem 1" in textes(dst_dir / "analyse" / "cas_analyse_74143_1.pptx")


def test_ranger_ids_dossier_prod_absent(tmp_path):
    root, _, _ = arbo(tmp_path)
    assert ranger_ids(str(root), "74143", None, ["74143_1"], "modifier", True, log) is None


def test_ranger_commune_copie_tout_et_vide_prod(tmp_path):
    root, src_dir, dst_dir = arbo(tmp_path)
    st = ranger_commune(str(tmp_path / "prod"), str(root), "74143", str(src_dir), True, log)
    assert st == {"excel": 1, "ppt": 1, "err": 0}
    assert load_workbook(dst_dir / "audit_74143.xlsx").active.max_row == 2   # Excel Prod IA (1 ligne)
    assert "Justification IA" in textes(dst_dir / "analyse" / "cas_analyse_74143_1.pptx")
    assert not src_dir.exists()


def test_ranger_commune_destination_absente(tmp_path):
    _, src_dir, _ = arbo(tmp_path)
    assert ranger_commune(str(tmp_path / "prod"), str(tmp_path / "rien"), "74143", str(src_dir), True, log) is None
    assert src_dir.exists()
