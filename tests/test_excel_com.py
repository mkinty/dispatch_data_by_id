"""Écriture via Excel (COM). Ignorés si Excel / pywin32 ne sont pas disponibles."""
import datetime

import pytest
from openpyxl import load_workbook

from dispatch_sna.services.excel_com import SessionExcel, excel_disponible
from tests.fichiers import excel

pytestmark = pytest.mark.skipif(not excel_disponible(), reason="Excel / pywin32 indisponible")

ENT = ["ID erreur", "Adresse", "Remarque", "Lien Street View", "Date Analyse", "Calcul", "Code"]


def _fichiers(tmp_path):
    dst = excel(tmp_path / "dst.xlsx", [
        ["74143_1", "adr SP", "vierge", None, None, "=LEN(B2)", "x"],
        ["74143_2", "adr SP 2", "garde", None, None, "=LEN(B3)", "y"],
    ], entetes=ENT, onglets=("Audit", "DGFIP"))
    src = excel(tmp_path / "src.xlsx", [
        ["74143_2", "adr IA 2", "Remarque IA", "https://sv/2?a=1&b=2", datetime.datetime(2026, 10, 16),
         "=LEN(B2)&\"!\"", "04121"],
    ], entetes=ENT)
    wb = load_workbook(src); ws = wb.active
    ws["D2"].hyperlink = "https://sv/2?a=1&b=2"
    ws["E2"].number_format = "dd/mm/yyyy"
    wb.save(src)
    return str(src), str(dst)


def test_ecriture_via_excel_et_verification(tmp_path):
    src, dst = _fichiers(tmp_path)
    with SessionExcel() as xl:
        r = xl.maj_lignes(src, dst, ["74143_2", "74143_9"],
                          ["Remarque", "Lien Street View", "Date Analyse", "Calcul", "Code"])
    assert r["maj"] == ["74143_2"] and r["non_verifies"] == [] and r["bloque"] is None
    assert r["absents_src"] == ["74143_9"]
    assert r["ouverture"] == "fichier local OneDrive"          # pas de montage SharePoint en test

    wb = load_workbook(dst); ws = wb["Audit"]
    assert [c.value for c in ws[2]][:3] == ["74143_1", "adr SP", "vierge"]       # autre ID intact
    assert ws["B3"].value == "adr SP 2"                                          # colonne non cochée intacte
    assert ws["C3"].value == "Remarque IA"
    assert ws["D3"].hyperlink.target == "https://sv/2?a=1&b=2"
    assert ws["E3"].value == datetime.datetime(2026, 10, 16) and ws["E3"].number_format == "dd/mm/yyyy"
    assert ws["F3"].value == '=LEN(B3)&"!"'                                       # formule recalée
    assert ws["G3"].value == "04121"                                              # reste du texte
    assert "DGFIP" in wb.sheetnames


def test_fichier_ouvert_ailleurs_est_signale_sans_rien_ecrire(tmp_path):
    src, dst = _fichiers(tmp_path)
    avant = open(dst, "rb").read()
    with SessionExcel() as collegue:                      # un « collègue » garde le classeur ouvert
        wb = collegue.xl.Workbooks.Open(dst)
        try:
            with SessionExcel() as xl:
                r = xl.maj_lignes(src, dst, ["74143_2"], ["Remarque"])
        finally:
            wb.Close(False)
    assert r["bloque"] and r["maj"] == []
    assert open(dst, "rb").read() == avant


def test_excel_occupe_l_appel_attend_au_lieu_d_echouer():
    """Reproduit « L'appel a été rejeté par l'appelé » : un autre client occupe Excel
    (gros recalcul) pendant que l'outil écrit. Grâce au filtre de messages, l'écriture
    attend qu'Excel soit libre puis réussit."""
    import threading, time
    import pythoncom, win32com.client as wc

    with SessionExcel() as s:
        wb = s.xl.Workbooks.Add()
        cible = wb.ActiveSheet.Range("L1")
        flux = pythoncom.CoMarshalInterThreadInterfaceInStream(pythoncom.IID_IDispatch, s.xl._oleobj_)
        fini = []

        def occupe():
            pythoncom.CoInitialize()
            x = wc.Dispatch(pythoncom.CoGetInterfaceAndReleaseStream(flux, pythoncom.IID_IDispatch))
            x.ActiveSheet.Range("A1:J150000").Formula = "=SQRT(ROW()*COLUMN())*RAND()"
            fini.append(time.time()); x = None
            pythoncom.CoUninitialize()

        th = threading.Thread(target=occupe); th.start(); time.sleep(1.0)
        cible.Value = "ecrit pendant l'occupation"   # sans filtre : com_error -2147418111
        ecrit = time.time()
        th.join()
        assert cible.Value == "ecrit pendant l'occupation"
        assert fini and ecrit >= fini[0] - 0.5         # l'écriture a bien attendu la fin du calcul
        cible = None; wb.Close(False); wb = None


def test_classeur_retente_si_excel_reste_occupe(tmp_path, monkeypatch):
    import pywintypes
    from dispatch_sna.services import excel_com
    src, dst = _fichiers(tmp_path)
    monkeypatch.setattr("time.sleep", lambda s: None)
    rejet = pywintypes.com_error(excel_com.RPC_E_CALL_REJECTED, "L'appel a été rejeté par l'appelé.", None, None)
    s = SessionExcel()
    appels = []

    def une_fois(lecture, dst, colonnes):
        appels.append(1)
        if len(appels) < 3:
            raise rejet
        return {"maj": ["74143_2"]}

    monkeypatch.setattr(s, "_maj_une_fois", une_fois)
    assert s.maj_lignes(src, dst, ["74143_2"], ["Remarque"]) == {"maj": ["74143_2"], "tentatives": 3}

    monkeypatch.setattr(s, "_maj_une_fois", lambda *a: (_ for _ in ()).throw(rejet))
    with pytest.raises(excel_com.ExcelOccupe, match="resté occupé"):
        s.maj_lignes(src, dst, ["74143_2"], ["Remarque"])


def test_autre_erreur_n_est_pas_retentee(tmp_path, monkeypatch):
    src, dst = _fichiers(tmp_path)
    s = SessionExcel()
    appels = []
    def une_fois(*a):
        appels.append(1); raise ValueError("colonne « ID erreur » absente")
    monkeypatch.setattr(s, "_maj_une_fois", une_fois)
    with pytest.raises(ValueError):
        s.maj_lignes(src, dst, ["74143_2"], ["Remarque"])
    assert len(appels) == 1


def test_aucune_colonne_n_ouvre_pas_le_classeur(tmp_path):
    src, dst = _fichiers(tmp_path)
    with SessionExcel() as xl:
        r = xl.maj_lignes(src, dst, ["74143_2"], [])
    assert r["maj"] == [] and r["ouverture"] is None
    assert r["donnees"]["74143_2"]["Remarque"] == "Remarque IA"
