import os
import re
import sys
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dispatch_ids import (extraire_ids, maj_lignes_excel, lire_lignes, lire_entetes,  # noqa: E402
                          modifier_ppt)
import dispatch_data  # noqa: E402

ENTETES = ["ID erreur", "Adresse", "Remarque", "Lien Street View", "Calcul"]


# ── Fixtures ──────────────────────────────────────────────────────
def _excel(path, lignes, entetes=ENTETES, onglets=("Audit",)):
    wb = Workbook(); ws = wb.active; ws.title = onglets[0]
    for autre in onglets[1:]:
        wb.create_sheet(autre)["A1"] = "garde-moi"
    ws.append(entetes)
    for l in lignes:
        ws.append(l)
    wb.save(path)
    return path


def _sp(nom, paras):
    ps = "".join(f"<a:p>{p}</a:p>" for p in paras)
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="1" name="{nom}"/></p:nvSpPr>'
            f"<p:txBody><a:bodyPr/><a:lstStyle/>{ps}</p:txBody></p:sp>")


def _run(t, extra=""):
    return f'<a:r><a:rPr lang="fr-FR" sz="1000">{extra}</a:rPr><a:t>{t}</a:t></a:r>'


def _pptx(path, shapes, rels=""):
    xml = ('<?xml version="1.0" encoding="UTF-8"?><p:sld xmlns:a="a" xmlns:p="p" xmlns:r="r">'
           f"<p:cSld><p:spTree>{''.join(shapes)}</p:spTree></p:cSld></p:sld>")
    rels = rels or '<?xml version="1.0"?><Relationships></Relationships>'
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/slides/slide1.xml", xml)
        z.writestr("ppt/slides/_rels/slide1.xml.rels", rels)
    return path


def _textes(path):
    xml = zipfile.ZipFile(path).read("ppt/slides/slide1.xml").decode()
    return [''.join(re.findall(r"<a:t[^>]*>(.*?)</a:t>", p, re.S))
            for p in re.findall(r"<a:p>.*?</a:p>", xml, re.S)]


def ppt_bp(path):
    """Structure du template bp_app (zone « Analyse » = Rectangle 8)."""
    return _pptx(path, [
        _sp("ZoneTexte 5", [_run("Analyses:")]),
        _sp("Rectangle 8", [_run("Ancienne analyse"), "<a:endParaRPr/>", ""]),
        _sp("Titre 2", [_run("74143 : adresse au 1 rue X (Etat IPE: OK)")]),
        _sp("ZoneTexte 3", [_run("SNA nombre de logement: 2")]),
        _sp("ZoneTexte 7", [_run("Lien Street View : ")]),
        _sp("TextBox 6", [_run("")]),
    ])


def ppt_babayaga(path):
    """Structure Babayaga : lien éclaté en plusieurs runs + run hyperlien."""
    hl = '<a:hlinkClick r:id="rId5"/>'
    rels = ('<?xml version="1.0"?><Relationships><Relationship Id="rId5" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
            'Target="https://old.example" TargetMode="External"/></Relationships>')
    return _pptx(path, [
        _sp("Titre 2", [_run("(2 - Maison): adresse au 3 rue Y (Etat IPE: OK)")]),
        _sp("ZoneTexte 4", [_run("IPE nombre de logement sur l'adresse: 1")]),
        _sp("ZoneTexte 5", [_run("Analyses:")]),
        _sp("Rectangle 8", [""]),
        _sp("ZoneTexte 6", [_run("   \t\t"), "", ""]),
        _sp("ZoneTexte 7", [_run("Justification IA")]),
        _sp("ZoneTexte 7", [_run("Lien Street ") + _run("View") + _run(" : ")
                            + _run("https://old.example", hl)]),
    ], rels)


# ── ID erreur ─────────────────────────────────────────────────────
def test_extraire_ids_groupe_et_dedoublonne():
    txt = "74143_1, 74143_2\n38068_207 74143_1 2a004_12 foo 123_4"
    assert extraire_ids(txt) == {"74143": ["74143_1", "74143_2"], "38068": ["38068_207"],
                                 "2A004": ["2A004_12"]}


# ── Excel ─────────────────────────────────────────────────────────
def test_maj_lignes_remplace_uniquement_les_ids(tmp_path):
    dst = _excel(tmp_path / "dst.xlsx", [
        ["74143_1", "A1", "vierge", None, "=LEN(B2)"],
        ["74143_2", "A2", "vierge", None, "=LEN(B3)"],
        ["74143_3", "A3", "garde", None, "=LEN(B4)"],
    ], onglets=("Audit", "Synthese"))
    # Source filtrée : lignes dans un autre ordre / à une autre position
    src = _excel(tmp_path / "src.xlsx", [
        ["74143_2", "A2", "Nouvelle remarque 2", "https://sv/2", "=LEN(B2)"],
        ["74143_3", "A3", "NE PAS COPIER", "https://sv/3", None],
    ])
    wb = load_workbook(src); ws = wb.active
    ws["D2"].hyperlink = "https://hl/2"
    ws["C2"].fill = PatternFill("solid", fgColor="FF00FF00")
    wb.save(src)

    r = maj_lignes_excel(str(src), str(dst), ["74143_2", "74143_9"])
    assert r["maj"] == ["74143_2"]
    assert r["absents_src"] == ["74143_9"]
    assert r["donnees"]["74143_2"] == {"Remarque": "Nouvelle remarque 2", "Lien Street View": "https://hl/2"}
    assert set(r["ids_src"]) == {"74143_2", "74143_3"}

    wb = load_workbook(dst); ws = wb["Audit"]
    assert [c.value for c in ws[2]][:3] == ["74143_1", "A1", "vierge"]
    assert [c.value for c in ws[3]][:4] == ["74143_2", "A2", "Nouvelle remarque 2", "https://sv/2"]
    assert ws["D3"].hyperlink.target == "https://hl/2"
    assert ws["C3"].fill.fgColor.rgb == "FF00FF00"
    assert ws["E3"].value == "=LEN(B3)"          # formule recalée sur la ligne cible
    assert ws["C4"].value == "garde"              # ID non demandé : intact
    assert wb["Synthese"]["A1"].value == "garde-moi"


def test_maj_lignes_id_absent_de_la_destination(tmp_path):
    dst = _excel(tmp_path / "dst.xlsx", [["74143_1", "A1", "x", None, None]])
    src = _excel(tmp_path / "src.xlsx", [["74143_5", "A5", "r", "l", None]])
    avant = os.path.getmtime(dst)
    r = maj_lignes_excel(str(src), str(dst), ["74143_5"])
    assert r["maj"] == [] and r["absents_dst"] == ["74143_5"]
    assert os.path.getmtime(dst) == avant           # rien à écrire → fichier non réécrit


def test_lire_lignes(tmp_path):
    src = _excel(tmp_path / "src.xlsx", [["74143_5", "A5", "rem", "lien", None]])
    assert lire_lignes(str(src), ["74143_5", "74143_6"]) == {
        "74143_5": {"Remarque": "rem", "Lien Street View": "lien"}}


# ── PPT ───────────────────────────────────────────────────────────
def test_modifier_ppt_template_bp(tmp_path):
    p = ppt_bp(tmp_path / "a.pptx")
    res = modifier_ppt(str(p), "Ligne 1 & <ok>\nLigne 2", "https://sv/new")
    assert res == {"analyse": True, "lien": True}
    t = _textes(p)
    assert "Ligne 1 &amp; &lt;ok&gt;" in t and "Ligne 2" in t
    assert "Ancienne analyse" not in t
    assert "Lien Street View : https://sv/new" in t
    assert "Analyses:" in t and "SNA nombre de logement: 2" in t
    zipfile.ZipFile(p).testzip()


def test_modifier_ppt_babayaga_hyperlien(tmp_path):
    p = ppt_babayaga(tmp_path / "b.pptx")
    res = modifier_ppt(str(p), "Nouvelle analyse", "https://sv/new?a=1&b=2")
    assert res == {"analyse": True, "lien": True}
    t = _textes(p)
    assert "Nouvelle analyse" in t and "Justification IA" not in t
    assert "Lien Street View : https://sv/new?a=1&amp;b=2" in t
    rels = zipfile.ZipFile(p).read("ppt/slides/_rels/slide1.xml.rels").decode()
    assert 'Target="https://sv/new?a=1&amp;b=2"' in rels
    assert "(2 - Maison): adresse au 3 rue Y (Etat IPE: OK)" in t


def test_modifier_ppt_analyse_vide_reprend_la_zone_avant_le_lien(tmp_path):
    p = _pptx(tmp_path / "c.pptx", [
        _sp("ZoneTexte 5", [_run("Analyses:")]),
        _sp("ZoneTexte 7", ['<a:endParaRPr lang="fr-FR" sz="1000"/>']),
        _sp("ZoneTexte 7", [_run("Lien Street View : x")]),
    ])
    assert modifier_ppt(str(p), "Remplie", None) == {"analyse": True, "lien": False}
    assert "Remplie" in _textes(p) and "Lien Street View : x" in _textes(p)


def test_modifier_ppt_none_ne_touche_pas(tmp_path):
    p = ppt_bp(tmp_path / "a.pptx")
    avant = open(p, "rb").read()
    assert modifier_ppt(str(p), None, None) == {"analyse": False, "lien": False}
    assert open(p, "rb").read() == avant


# ── ranger_ids (bout en bout) ─────────────────────────────────────
def _arbo(tmp_path):
    root = tmp_path / "sp"; prod = tmp_path / "prod"
    dst_dir = root / "Dep74" / "74143"; (dst_dir / "analyse").mkdir(parents=True)
    src_dir = prod / "74143"; src_dir.mkdir(parents=True)
    _excel(dst_dir / "audit_74143.xlsx", [["74143_1", "A1", "vierge", None, None],
                                          ["74143_2", "A2", "vierge", None, None]])
    _excel(src_dir / "audit_74143.xlsx", [["74143_1", "A1", "Rem 1", "https://sv/1", None]])
    ppt_bp(dst_dir / "analyse" / "cas_analyse_74143_1.pptx")
    ppt_babayaga(src_dir / "cas_analyse_74143_1.pptx")
    return root, src_dir, dst_dir


def _log(*a):
    pass


def test_ranger_ids_modifier(tmp_path):
    root, src_dir, dst_dir = _arbo(tmp_path)
    st = dispatch_data.ranger_ids(str(root), "74143", str(src_dir), ["74143_1"], "modifier", True, _log)
    assert st == {"lignes": 1, "excel": 1, "ppt": 1, "err": 0}
    t = _textes(dst_dir / "analyse" / "cas_analyse_74143_1.pptx")
    assert "Rem 1" in t and "Lien Street View : https://sv/1" in t
    assert "SNA nombre de logement: 2" in t        # c'est bien l'ancien PPT modifié
    assert not src_dir.exists()                     # tout traité → dossier supprimé


def test_ranger_ids_remplacer_et_conserve_dossier_si_autres_ids(tmp_path):
    root, src_dir, dst_dir = _arbo(tmp_path)
    _excel(src_dir / "audit_74143.xlsx", [["74143_1", "A1", "Rem 1", "u", None],
                                          ["74143_2", "A2", "Rem 2", "u", None]])
    st = dispatch_data.ranger_ids(str(root), "74143", str(src_dir), ["74143_1"], "remplacer", True, _log)
    assert st["err"] == 0 and st["ppt"] == 1
    assert "Justification IA" in _textes(dst_dir / "analyse" / "cas_analyse_74143_1.pptx")
    assert src_dir.exists()                         # 74143_2 non traité → gardé
    ws = load_workbook(dst_dir / "audit_74143.xlsx").active
    assert ws["C3"].value == "vierge"


def test_ranger_ids_sans_option_ppt(tmp_path):
    root, src_dir, dst_dir = _arbo(tmp_path)
    avant = open(dst_dir / "analyse" / "cas_analyse_74143_1.pptx", "rb").read()
    st = dispatch_data.ranger_ids(str(root), "74143", str(src_dir), ["74143_1"], None, False, _log)
    assert st == {"lignes": 1, "excel": 1, "ppt": 0, "err": 0}
    assert open(dst_dir / "analyse" / "cas_analyse_74143_1.pptx", "rb").read() == avant


def test_ranger_ids_dossier_prod_absent(tmp_path):
    root, _, _ = _arbo(tmp_path)
    assert dispatch_data.ranger_ids(str(root), "74143", None, ["74143_1"], "modifier", True, _log) is None


# ── Colonnes à mettre à jour ──────────────────────────────────────
def test_maj_lignes_colonnes_cochees_uniquement(tmp_path):
    dst = _excel(tmp_path / "dst.xlsx", [["74143_1", "adr SP", "vierge", "lien SP", "calc SP"]])
    src = _excel(tmp_path / "src.xlsx", [["74143_1", "adr IA", "rem IA", "lien IA", "calc IA"]])
    r = maj_lignes_excel(str(src), str(dst), ["74143_1"], ["Remarque", "Inconnue"])
    assert r["maj"] == ["74143_1"] and r["col_absentes"] == ["Inconnue"]
    ws = load_workbook(dst).active
    assert [c.value for c in ws[2]] == ["74143_1", "adr SP", "rem IA", "lien SP", "calc SP"]


def test_maj_lignes_aucune_colonne_ne_modifie_pas(tmp_path):
    dst = _excel(tmp_path / "dst.xlsx", [["74143_1", "adr SP", "vierge", None, None]])
    src = _excel(tmp_path / "src.xlsx", [["74143_1", "adr IA", "rem IA", "lien IA", None]])
    avant = open(dst, "rb").read()
    r = maj_lignes_excel(str(src), str(dst), ["74143_1"], [])
    assert r["maj"] == [] and r["absents_dst"] == []
    assert r["donnees"]["74143_1"]["Remarque"] == "rem IA"     # dispo pour les PPT
    assert open(dst, "rb").read() == avant


def test_ranger_ids_aucune_colonne_mais_ppt_modifie(tmp_path):
    root, src_dir, dst_dir = _arbo(tmp_path)
    st = dispatch_data.ranger_ids(str(root), "74143", str(src_dir), ["74143_1"], "modifier", True, _log,
                                  colonnes=[])
    assert st == {"lignes": 0, "excel": 0, "ppt": 1, "err": 0}
    assert load_workbook(dst_dir / "audit_74143.xlsx").active["C2"].value == "vierge"
    assert "Rem 1" in _textes(dst_dir / "analyse" / "cas_analyse_74143_1.pptx")


def test_lire_entetes_et_detection(tmp_path):
    root, src_dir, _ = _arbo(tmp_path)
    assert lire_entetes(str(src_dir / "audit_74143.xlsx")) == ENTETES
    cols, f = dispatch_data.detecter_colonnes(str(tmp_path / "prod"), str(root))
    assert cols == ENTETES and f.endswith("audit_74143.xlsx") and "prod" in f
    # Prod IA vide → repli sur l'Excel SharePoint
    cols, f = dispatch_data.detecter_colonnes(str(tmp_path / "vide"), str(root))
    assert cols == ENTETES and os.path.join("Dep74", "74143") in f
    assert dispatch_data.detecter_colonnes(str(tmp_path / "x"), str(tmp_path / "y")) == ([], None)
