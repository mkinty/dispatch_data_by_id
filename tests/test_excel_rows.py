import os

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from dispatch_sna.services.excel_rows import maj_lignes_excel, lire_lignes, lire_entetes
from tests.fichiers import ENTETES, excel


def test_maj_lignes_remplace_uniquement_les_ids(tmp_path):
    dst = excel(tmp_path / "dst.xlsx", [
        ["74143_1", "A1", "vierge", None, "=LEN(B2)"],
        ["74143_2", "A2", "vierge", None, "=LEN(B3)"],
        ["74143_3", "A3", "garde", None, "=LEN(B4)"],
    ], onglets=("Audit", "Synthese"))
    # Source filtrée : lignes dans un autre ordre / à une autre position
    src = excel(tmp_path / "src.xlsx", [
        ["74143_2", "A2", "Nouvelle remarque 2", "https://sv/2", "=LEN(B2)"],
        ["74143_3", "A3", "NE PAS COPIER", "https://sv/3", None],
    ])
    wb = load_workbook(src); ws = wb.active
    ws["D2"].hyperlink = "https://hl/2"
    ws["C2"].fill = PatternFill("solid", fgColor="FF00FF00")
    wb.save(src)

    r = maj_lignes_excel(str(src), str(dst), ["74143_2", "74143_9"])
    assert r["maj"] == ["74143_2"]
    assert r["absents_src"] == ["74143_9"]
    assert r["donnees"]["74143_2"] == {"Remarque": "Nouvelle remarque 2", "Lien Street View": "https://hl/2"}
    assert set(r["ids_src"]) == {"74143_2", "74143_3"}

    wb = load_workbook(dst); ws = wb["Audit"]
    assert [c.value for c in ws[2]][:3] == ["74143_1", "A1", "vierge"]
    assert [c.value for c in ws[3]][:4] == ["74143_2", "A2", "Nouvelle remarque 2", "https://sv/2"]
    assert ws["D3"].hyperlink.target == "https://hl/2"
    assert ws["C3"].fill.fgColor.rgb == "FF00FF00"
    assert ws["E3"].value == "=LEN(B3)"          # formule recalée sur la ligne cible
    assert ws["C4"].value == "garde"              # ID non demandé : intact
    assert wb["Synthese"]["A1"].value == "garde-moi"


def test_maj_lignes_id_absent_de_la_destination(tmp_path):
    dst = excel(tmp_path / "dst.xlsx", [["74143_1", "A1", "x", None, None]])
    src = excel(tmp_path / "src.xlsx", [["74143_5", "A5", "r", "l", None]])
    avant = os.path.getmtime(dst)
    r = maj_lignes_excel(str(src), str(dst), ["74143_5"])
    assert r["maj"] == [] and r["absents_dst"] == ["74143_5"]
    assert os.path.getmtime(dst) == avant           # rien à écrire → fichier non réécrit


def test_maj_lignes_colonnes_cochees_uniquement(tmp_path):
    dst = excel(tmp_path / "dst.xlsx", [["74143_1", "adr SP", "vierge", "lien SP", "calc SP"]])
    src = excel(tmp_path / "src.xlsx", [["74143_1", "adr IA", "rem IA", "lien IA", "calc IA"]])
    r = maj_lignes_excel(str(src), str(dst), ["74143_1"], ["Remarque", "Inconnue"])
    assert r["maj"] == ["74143_1"] and r["col_absentes"] == ["Inconnue"]
    ws = load_workbook(dst).active
    assert [c.value for c in ws[2]] == ["74143_1", "adr SP", "rem IA", "lien SP", "calc SP"]


def test_maj_lignes_aucune_colonne_ne_modifie_pas(tmp_path):
    dst = excel(tmp_path / "dst.xlsx", [["74143_1", "adr SP", "vierge", None, None]])
    src = excel(tmp_path / "src.xlsx", [["74143_1", "adr IA", "rem IA", "lien IA", None]])
    avant = open(dst, "rb").read()
    r = maj_lignes_excel(str(src), str(dst), ["74143_1"], [])
    assert r["maj"] == [] and r["absents_dst"] == []
    assert r["donnees"]["74143_1"]["Remarque"] == "rem IA"     # dispo pour les PPT
    assert open(dst, "rb").read() == avant


def test_lire_lignes(tmp_path):
    src = excel(tmp_path / "src.xlsx", [["74143_5", "A5", "rem", "lien", None]])
    assert lire_lignes(str(src), ["74143_5", "74143_6"]) == {
        "74143_5": {"Remarque": "rem", "Lien Street View": "lien"}}


def test_lire_entetes(tmp_path):
    src = excel(tmp_path / "src.xlsx", [])
    assert lire_entetes(str(src)) == ENTETES
