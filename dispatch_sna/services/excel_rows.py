"""Excel audit : lecture des en-têtes / lignes, remplacement des lignes par ID erreur."""
from __future__ import annotations

import os
from copy import copy

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from .constantes import ONGLET_AUDIT, COL_ID_ERR, COL_REMARQUE, COL_LIEN_SV


def _norm(v) -> str:
    return str(v).strip().upper() if v is not None else ""


def trouver_onglet(wb, nom=ONGLET_AUDIT):
    """Onglet « Audit » (espaces/casse ignorés), repli sur l'onglet actif."""
    cible = "".join(nom.split()).casefold()
    for titre in wb.sheetnames:
        if "".join(titre.split()).casefold() == cible:
            return wb[titre]
    return wb.active


def _entetes(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def _index_ids(ws, col_id: int) -> dict[str, int]:
    """{ID erreur → n° de ligne} (première occurrence)."""
    idx = {}
    for r in range(2, ws.max_row + 1):
        v = _norm(ws.cell(r, col_id).value)
        if v and v not in idx:
            idx[v] = r
    return idx


def _valeur_cellule(cell) -> str:
    """Texte d'une cellule — l'hyperlien est préféré au texte affiché."""
    hl = getattr(cell, "hyperlink", None)
    cible = getattr(hl, "target", None) if hl is not None else None
    if cible:
        return str(cible).strip()
    return "" if cell.value is None else str(cell.value).strip()


def _copier_cellule(src, dst):
    v = src.value
    if isinstance(v, str) and v.startswith("=") and src.coordinate != dst.coordinate:
        # Recale les références relatives sur la ligne de destination
        try: v = Translator(v, origin=src.coordinate).translate_formula(dst.coordinate)
        except Exception: pass
    dst.value = v
    dst.hyperlink = copy(src.hyperlink) if src.hyperlink else None
    if src.has_style:
        dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
        dst.alignment = copy(src.alignment); dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def lire_entetes(chemin: str) -> list[str]:
    """En-têtes (ligne 1) de l'onglet Audit, dans l'ordre du fichier."""
    wb = load_workbook(chemin, read_only=True)
    try:
        ws = trouver_onglet(wb)
        ligne = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return list(dict.fromkeys(str(v).strip() for v in ligne if v is not None and str(v).strip()))
    finally:
        wb.close()


def lire_lignes(chemin: str, ids) -> dict[str, dict[str, str]]:
    """{ID → {"Remarque": ..., "Lien Street View": ...}} pour les ID trouvés."""
    wb = load_workbook(chemin)
    ws = trouver_onglet(wb)
    h = _entetes(ws)
    if COL_ID_ERR not in h:
        return {}
    idx = _index_ids(ws, h[COL_ID_ERR])
    res = {}
    for i in ids:
        r = idx.get(_norm(i))
        if r is None:
            continue
        res[_norm(i)] = {col: (_valeur_cellule(ws.cell(r, h[col])) if col in h else "")
                         for col in (COL_REMARQUE, COL_LIEN_SV)}
    return res


class ExcelErreur(Exception):
    """Structure inattendue (colonne ID erreur absente...)."""


def maj_lignes_excel(src: str, dst: str, ids, colonnes=None) -> dict:
    """Remplace dans dst les lignes des ID donnés par celles de src.

    Colonnes associées par nom d'en-tête ; valeurs, hyperliens et mise en forme
    copiés. `colonnes` : noms des colonnes à mettre à jour (None = toutes les
    colonnes communes, [] = aucune → dst non modifié). Les autres lignes et
    colonnes de dst ne sont pas touchées. Retourne
    {"maj": [...], "absents_src": [...], "absents_dst": [...],
     "donnees": {ID → {Remarque, Lien Street View}}, "ids_src": [...],
     "col_absentes": [...]}.
    Lève PermissionError si dst est verrouillé.
    """
    wanted = list(dict.fromkeys(_norm(i) for i in ids))
    wb_s = load_workbook(src); ws_s = trouver_onglet(wb_s)
    wb_d = load_workbook(dst); ws_d = trouver_onglet(wb_d)
    h_s, h_d = _entetes(ws_s), _entetes(ws_d)
    for nom, h in (("Prod IA", h_s), ("SharePoint", h_d)):
        if COL_ID_ERR not in h:
            raise ExcelErreur(f"colonne « {COL_ID_ERR} » absente de l'Excel {nom}")
    idx_s = _index_ids(ws_s, h_s[COL_ID_ERR])
    idx_d = _index_ids(ws_d, h_d[COL_ID_ERR])
    voulues = None if colonnes is None else {str(c).strip() for c in colonnes} - {COL_ID_ERR}
    paires = [(h_s[n], c) for n, c in h_d.items()
              if n in h_s and n != COL_ID_ERR and (voulues is None or n in voulues)]

    res = {"maj": [], "absents_src": [], "absents_dst": [], "donnees": {},
           "ids_src": list(idx_s),
           "col_absentes": sorted(voulues - (h_s.keys() & h_d.keys())) if voulues else []}
    for i in wanted:
        rs = idx_s.get(i)
        if rs is None:
            res["absents_src"].append(i); continue
        res["donnees"][i] = {col: (_valeur_cellule(ws_s.cell(rs, h_s[col])) if col in h_s else "")
                             for col in (COL_REMARQUE, COL_LIEN_SV)}
        if not paires:
            continue
        rd = idx_d.get(i)
        if rd is None:
            res["absents_dst"].append(i); continue
        for cs, cd in paires:
            _copier_cellule(ws_s.cell(rs, cs), ws_d.cell(rd, cd))
        res["maj"].append(i)

    if res["maj"]:
        tmp = dst + ".tmp.xlsx"
        try:
            wb_d.save(tmp)
            os.replace(tmp, dst)
        finally:
            if os.path.isfile(tmp):
                os.remove(tmp)
    return res
