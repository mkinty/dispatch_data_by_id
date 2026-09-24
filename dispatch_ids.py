"""
Mise à jour ciblée par ID erreur (aucune dépendance à l'interface).

- extraire_ids        : ID erreur collés par l'utilisateur, regroupés par commune
- maj_lignes_excel    : remplace, dans l'Excel SharePoint, uniquement les lignes
                        des ID erreur donnés par celles de l'Excel Prod IA
- modifier_ppt        : réécrit les zones « Analyse » et « Lien Street View »
                        d'un PPT existant (édition XML directe, sans python-pptx)
"""
from __future__ import annotations

import os
import re
import shutil
import zipfile
from copy import copy
from html import unescape

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

ONGLET_AUDIT = "Audit"
COL_ID_ERR = "ID erreur"
COL_REMARQUE = "Remarque"
COL_LIEN_SV = "Lien Street View"
PPT_PREFIX = "cas_analyse_"

# <code INSEE>_<numéro>, ex : 74143_1, 2A004_12 (même format que export_lot_by_id)
ERROR_ID_PATTERN = re.compile(r'\b((?:2[AB]|\d{2})\d{3})_(\d+)\b')


# ── ID erreur ─────────────────────────────────────────────────────
def extraire_ids(texte: str) -> dict[str, list[str]]:
    """{"74143": ["74143_1", "74143_2"], ...} — ordre conservé, doublons supprimés."""
    groupes: dict[str, list[str]] = {}
    for insee, num in ERROR_ID_PATTERN.findall(texte.upper()):
        ids = groupes.setdefault(insee, [])
        id_err = f"{insee}_{num}"
        if id_err not in ids:
            ids.append(id_err)
    return groupes


def nom_ppt(id_err: str) -> str:
    return f"{PPT_PREFIX}{id_err}.pptx"


# ── Excel ─────────────────────────────────────────────────────────
def _norm(v) -> str:
    return str(v).strip().upper() if v is not None else ""


def trouver_onglet(wb, nom=ONGLET_AUDIT):
    """Onglet « Audit » (espaces/casse ignorés), repli sur l'onglet actif."""
    cible = "".join(nom.split()).casefold()
    for titre in wb.sheetnames:
        if "".join(titre.split()).casefold() == cible:
            return wb[titre]
    return wb.active


def _entetes(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def _index_ids(ws, col_id: int) -> dict[str, int]:
    """{ID erreur → n° de ligne} (première occurrence)."""
    idx = {}
    for r in range(2, ws.max_row + 1):
        v = _norm(ws.cell(r, col_id).value)
        if v and v not in idx:
            idx[v] = r
    return idx


def _valeur_cellule(cell) -> str:
    """Texte d'une cellule — l'hyperlien est préféré au texte affiché."""
    hl = getattr(cell, "hyperlink", None)
    cible = getattr(hl, "target", None) if hl is not None else None
    if cible:
        return str(cible).strip()
    return "" if cell.value is None else str(cell.value).strip()


def _copier_cellule(src, dst):
    v = src.value
    if isinstance(v, str) and v.startswith("=") and src.coordinate != dst.coordinate:
        # Recale les références relatives sur la ligne de destination
        try: v = Translator(v, origin=src.coordinate).translate_formula(dst.coordinate)
        except Exception: pass
    dst.value = v
    dst.hyperlink = copy(src.hyperlink) if src.hyperlink else None
    if src.has_style:
        dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
        dst.alignment = copy(src.alignment); dst.number_format = src.number_format
        dst.protection = copy(src.protection)


def lire_entetes(chemin: str) -> list[str]:
    """En-têtes (ligne 1) de l'onglet Audit, dans l'ordre du fichier."""
    wb = load_workbook(chemin, read_only=True)
    try:
        ws = trouver_onglet(wb)
        ligne = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return list(dict.fromkeys(str(v).strip() for v in ligne if v is not None and str(v).strip()))
    finally:
        wb.close()


def lire_lignes(chemin: str, ids) -> dict[str, dict[str, str]]:
    """{ID → {"Remarque": ..., "Lien Street View": ...}} pour les ID trouvés."""
    wb = load_workbook(chemin)
    ws = trouver_onglet(wb)
    h = _entetes(ws)
    if COL_ID_ERR not in h:
        return {}
    idx = _index_ids(ws, h[COL_ID_ERR])
    res = {}
    for i in ids:
        r = idx.get(_norm(i))
        if r is None:
            continue
        res[_norm(i)] = {col: (_valeur_cellule(ws.cell(r, h[col])) if col in h else "")
                         for col in (COL_REMARQUE, COL_LIEN_SV)}
    return res


class ExcelErreur(Exception):
    """Structure inattendue (colonne ID erreur absente...)."""


def maj_lignes_excel(src: str, dst: str, ids, colonnes=None) -> dict:
    """Remplace dans dst les lignes des ID donnés par celles de src.

    Colonnes associées par nom d'en-tête ; valeurs, hyperliens et mise en forme
    copiés. `colonnes` : noms des colonnes à mettre à jour (None = toutes les
    colonnes communes, [] = aucune → dst non modifié). Les autres lignes et
    colonnes de dst ne sont pas touchées. Retourne
    {"maj": [...], "absents_src": [...], "absents_dst": [...],
     "donnees": {ID → {Remarque, Lien Street View}}, "ids_src": [...],
     "col_absentes": [...]}.
    Lève PermissionError si dst est verrouillé.
    """
    wanted = list(dict.fromkeys(_norm(i) for i in ids))
    wb_s = load_workbook(src); ws_s = trouver_onglet(wb_s)
    wb_d = load_workbook(dst); ws_d = trouver_onglet(wb_d)
    h_s, h_d = _entetes(ws_s), _entetes(ws_d)
    for nom, h in (("Prod IA", h_s), ("SharePoint", h_d)):
        if COL_ID_ERR not in h:
            raise ExcelErreur(f"colonne « {COL_ID_ERR} » absente de l'Excel {nom}")
    idx_s = _index_ids(ws_s, h_s[COL_ID_ERR])
    idx_d = _index_ids(ws_d, h_d[COL_ID_ERR])
    voulues = None if colonnes is None else {str(c).strip() for c in colonnes} - {COL_ID_ERR}
    paires = [(h_s[n], c) for n, c in h_d.items()
              if n in h_s and n != COL_ID_ERR and (voulues is None or n in voulues)]

    res = {"maj": [], "absents_src": [], "absents_dst": [], "donnees": {},
           "ids_src": list(idx_s),
           "col_absentes": sorted(voulues - (h_s.keys() & h_d.keys())) if voulues else []}
    for i in wanted:
        rs = idx_s.get(i)
        if rs is None:
            res["absents_src"].append(i); continue
        res["donnees"][i] = {col: (_valeur_cellule(ws_s.cell(rs, h_s[col])) if col in h_s else "")
                             for col in (COL_REMARQUE, COL_LIEN_SV)}
        if not paires:
            continue
        rd = idx_d.get(i)
        if rd is None:
            res["absents_dst"].append(i); continue
        for cs, cd in paires:
            _copier_cellule(ws_s.cell(rs, cs), ws_d.cell(rd, cd))
        res["maj"].append(i)

    if res["maj"]:
        tmp = dst + ".tmp.xlsx"
        try:
            wb_d.save(tmp)
            os.replace(tmp, dst)
        finally:
            if os.path.isfile(tmp):
                os.remove(tmp)
    return res


# ── PPT ───────────────────────────────────────────────────────────
_SHAPE = re.compile(r"<p:sp>.*?</p:sp>", re.S)
_TXBODY = re.compile(r"(<p:txBody>)(.*?)(</p:txBody>)", re.S)
_PARA = re.compile(r"<a:p(?:\s[^>]*)?>.*?</a:p>", re.S)   # n'attrape pas <a:p/> ni <a:pPr>
_RUN = re.compile(r"<a:r>.*?</a:r>", re.S)
_T = re.compile(r"(<a:t(?:\s[^>]*)?>)(.*?)(</a:t>)", re.S)
_END_RPR = re.compile(r"<a:endParaRPr\b[^>]*?(?:/>|>.*?</a:endParaRPr>)", re.S)
_HLINK = re.compile(r'<a:hlinkClick\b[^>]*\br:id="([^"]+)"')

# Libellés fixes du template : jamais la zone « Analyse »
_LIBELLES = re.compile(r"^(analyses\s*:|lien street view|sna nombre|ipe nombre|type de probl"
                       r"|\{\{photo|photo$)", re.I)
_LABEL_LIEN = "Lien Street View : "


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def _texte(xml: str) -> str:
    return unescape("".join(m.group(2) for m in _T.finditer(xml)))


def _est_libelle(texte: str) -> bool:
    t = texte.strip()
    return bool(_LIBELLES.match(t)) or "adresse au" in t.lower()


def _set_para(para: str, texte: str) -> str:
    """Met `texte` dans le 1er fragment <a:t> du paragraphe, vide les autres."""
    if _T.search(para):
        n = {"i": 0}
        def rep(m):
            n["i"] += 1
            return f"{m.group(1)}{_esc(texte) if n['i'] == 1 else ''}{m.group(3)}"
        return _T.sub(rep, para)
    # Paragraphe vide : on crée un run avec la mise en forme de fin de paragraphe
    m = _END_RPR.search(para)
    rpr = m.group(0).replace("a:endParaRPr", "a:rPr") if m else ""
    run = f"<a:r>{rpr}<a:t>{_esc(texte)}</a:t></a:r>"
    pos = m.start() if m else para.rindex("</a:p>")
    return para[:pos] + run + para[pos:]


def _remplir_analyse(sp: str, texte: str) -> str:
    """Remplace le texte d'une zone : 1ère ligne dans le 1er paragraphe écrit,
    lignes suivantes en paragraphes clonés ; anciens paragraphes texte retirés,
    paragraphes vides (espacement) conservés."""
    mb = _TXBODY.search(sp)
    if not mb:
        return sp
    body = mb.group(2)
    paras = list(_PARA.finditer(body))
    if not paras:
        return sp
    avec_texte = [p for p in paras if _texte(p.group(0)).strip()]
    modele = (avec_texte or paras)[0]
    lignes = texte.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    nouveau = "".join(_set_para(modele.group(0), l) for l in lignes)
    out, pos = [], 0
    for p in paras:
        out.append(body[pos:p.start()])
        if p is modele:
            out.append(nouveau)
        elif p not in avec_texte:
            out.append(p.group(0))
        pos = p.end()
    out.append(body[pos:])
    body2 = "".join(out)
    return sp[:mb.start(2)] + body2 + sp[mb.end(2):]


def _remplir_lien(para: str, lien: str, rels: str) -> tuple[str, str]:
    """Réécrit « Lien Street View : <lien> ». Si un run porte déjà un
    hyperlien, le libellé va dans le 1er run et l'URL dans ce run (cible mise à jour)."""
    runs = list(_RUN.finditer(para))
    lien_run = next((r for r in runs[1:] if _HLINK.search(r.group(0))), None)
    if lien_run is None:
        m = _HLINK.search(runs[0].group(0)) if runs else None
        if m:
            rels = _maj_rel(rels, m.group(1), lien)
        return _set_para(para, f"{_LABEL_LIEN}{lien}".rstrip()), rels
    rels = _maj_rel(rels, _HLINK.search(lien_run.group(0)).group(1), lien)
    out, pos = [], 0
    for r in runs:
        txt = _LABEL_LIEN if r is runs[0] else (lien if r is lien_run else "")
        out.append(para[pos:r.start()])
        out.append(_T.sub(lambda m: f"{m.group(1)}{_esc(txt)}{m.group(3)}", r.group(0), count=1))
        pos = r.end()
    out.append(para[pos:])
    return "".join(out), rels


def _maj_rel(rels: str, rid: str, cible: str) -> str:
    if not cible:
        return rels
    return re.sub(rf'(<Relationship\b[^>]*\bId="{re.escape(rid)}"[^>]*\bTarget=")[^"]*(")',
                  lambda m: f"{m.group(1)}{_esc(cible)}{m.group(2)}", rels, count=1)


def _modifier_slide(xml: str, rels: str, analyse: str | None, lien: str | None):
    """Retourne (xml, rels, analyse_ok, lien_ok)."""
    shapes = list(_SHAPE.finditer(xml))
    textes = [_texte(s.group(0)) for s in shapes]
    a_ok = l_ok = False
    remplacements = {}   # index shape → nouveau XML

    # Lien Street View : paragraphe commençant par « Lien Street View »
    i_lien = None
    if lien is not None:
        for i, s in enumerate(shapes):
            sp = s.group(0); changed = False
            for p in list(_PARA.finditer(sp))[::-1]:
                if _texte(p.group(0)).strip().lower().startswith("lien street view"):
                    np_, rels = _remplir_lien(p.group(0), lien, rels)
                    sp = sp[:p.start()] + np_ + sp[p.end():]; changed = True
            if changed:
                remplacements[i] = sp; l_ok = True
                if i_lien is None: i_lien = i
    else:
        i_lien = next((i for i, t in enumerate(textes)
                       if t.strip().lower().startswith("lien street view")), None)

    # Analyse : uniquement sur la slide qui porte le libellé « Analyses: »
    if analyse is not None and any(t.strip().lower().startswith("analyses") for t in textes):
        cible = next((i for i, t in enumerate(textes) if t.strip() and not _est_libelle(t)), None)
        if cible is None and i_lien is not None:   # zone vide : celle juste avant le lien
            j = i_lien - 1
            if j >= 0 and not _est_libelle(textes[j]) and _TXBODY.search(shapes[j].group(0)):
                cible = j
        if cible is not None:
            remplacements[cible] = _remplir_analyse(remplacements.get(cible, shapes[cible].group(0)), analyse)
            a_ok = True

    if remplacements:
        out, pos = [], 0
        for i, s in enumerate(shapes):
            out.append(xml[pos:s.start()]); out.append(remplacements.get(i, s.group(0))); pos = s.end()
        out.append(xml[pos:]); xml = "".join(out)
    return xml, rels, a_ok, l_ok


def modifier_ppt(chemin: str, analyse: str | None, lien: str | None) -> dict:
    """Réécrit Analyse et/ou Lien Street View dans le PPT (None = ne pas toucher).

    Retourne {"analyse": bool, "lien": bool} — ce qui a effectivement été trouvé
    et modifié. Écriture via fichier temporaire. Lève PermissionError si verrouillé.
    """
    slide_re = re.compile(r"^ppt/slides/slide\d+\.xml$")
    nouveaux, res = {}, {"analyse": False, "lien": False}
    with zipfile.ZipFile(chemin) as z:
        noms = z.namelist()
        for nom in sorted(n for n in noms if slide_re.match(n)):
            rel_nom = nom.replace("slides/", "slides/_rels/") + ".rels"
            xml = z.read(nom).decode("utf-8")
            rels = z.read(rel_nom).decode("utf-8") if rel_nom in noms else ""
            xml2, rels2, a_ok, l_ok = _modifier_slide(
                xml, rels, None if res["analyse"] else analyse, lien)
            res["analyse"] |= a_ok; res["lien"] |= l_ok
            if xml2 != xml: nouveaux[nom] = xml2
            if rels2 != rels: nouveaux[rel_nom] = rels2
    if not nouveaux:
        return res
    tmp = chemin + ".tmp.pptx"
    try:
        with zipfile.ZipFile(chemin) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
            for item in zi.infolist():
                data = nouveaux[item.filename].encode("utf-8") if item.filename in nouveaux \
                    else zi.read(item.filename)
                zo.writestr(item, data)
        os.replace(tmp, chemin)
    finally:
        if os.path.isfile(tmp):
            os.remove(tmp)
    return res


def copier_ppt(src: str, dst: str):
    tmp = dst + ".tmp"
    shutil.copy2(src, tmp)
    os.replace(tmp, dst)
